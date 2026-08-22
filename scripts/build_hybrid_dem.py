#!/usr/bin/env python3
"""Build a robust FABDEM/station hybrid DEM for the Cameroon gravity project.

The 30 m FABDEM morphology is preserved. Only a smooth, cross-validated
station-minus-FABDEM residual correction is added. The program uses the exact
input measurements without editing them and reports every rejected control.
"""

from __future__ import annotations

import argparse
import json
import math
from collections import defaultdict
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from PIL import Image, TiffImagePlugin
from scipy.interpolate import RegularGridInterpolator
from scipy.ndimage import zoom
from scipy.spatial import Delaunay, cKDTree


RESOLUTION_DEG = 1.0 / 3600.0
COARSE_FACTOR = 15
COARSE_SIZE = 3600 // COARSE_FACTOR
COARSE_PADDING = 4
LAT_KM = 110.574
LON_KM = 111.320
IDW_NEIGHBORS = 32
IDW_POWER = 1.5
IDW_SOFTENING_KM = 5.0
RESIDUAL_CLIP_M = 40.0
OUTLIER_THRESHOLD_M = 150.0
OUTSIDE_TAPER_KM = 20.0
BLOCK_SIZE_KM = 25.0
CV_FOLDS = 5
RANDOM_SEED = 20260822


def load_stations(path: Path) -> np.ndarray:
    ws = load_workbook(path, read_only=True, data_only=True)["Feuil1"]
    return np.asarray(
        [[float(v) for v in row[:4]] for row in ws.iter_rows(min_row=2, values_only=True)],
        dtype=float,
    )


def fabdem_path(root: Path, lat0: int, lon0: int) -> Path:
    return root / f"N{lat0:02d}E{lon0:03d}_FABDEM_V1-2.tif"


def copernicus_path(root: Path, lat0: int, lon0: int) -> Path:
    stem = f"Copernicus_DSM_COG_10_N{lat0:02d}_00_E{lon0:03d}_00_DEM"
    return root / f"{stem}.tif"


def bilinear_from_array(array: np.ndarray, lon: float, lat: float, lon0: int, lat0: int) -> float:
    x = (lon - lon0) / RESOLUTION_DEG - 0.5
    y = ((lat0 + 1) - lat) / RESOLUTION_DEG - 0.5
    x_floor = math.floor(x)
    y_floor = math.floor(y)
    fx = x - x_floor
    fy = y - y_floor
    ix = max(0, min(array.shape[1] - 2, int(x_floor)))
    iy = max(0, min(array.shape[0] - 2, int(y_floor)))
    return float(
        (1 - fy) * ((1 - fx) * array[iy, ix] + fx * array[iy, ix + 1])
        + fy * ((1 - fx) * array[iy + 1, ix] + fx * array[iy + 1, ix + 1])
    )


def sample_tiles(stations: np.ndarray, root: Path, product: str) -> np.ndarray:
    result = np.full(len(stations), np.nan, dtype=float)
    groups: dict[tuple[int, int], list[int]] = defaultdict(list)
    for i, (lon, lat) in enumerate(stations[:, :2]):
        groups[(math.floor(lat), math.floor(lon))].append(i)

    for (lat0, lon0), indices in groups.items():
        path = fabdem_path(root, lat0, lon0) if product == "fabdem" else copernicus_path(root, lat0, lon0)
        if not path.exists():
            raise FileNotFoundError(f"Missing {product} tile: {path}")
        with Image.open(path) as image:
            raster = np.asarray(image, dtype=np.float32)
        for i in indices:
            result[i] = bilinear_from_array(raster, stations[i, 0], stations[i, 1], lon0, lat0)
    return result


def to_xy_km(lon_lat: np.ndarray, latitude_reference: float) -> np.ndarray:
    return np.column_stack(
        [lon_lat[:, 0] * LON_KM * math.cos(math.radians(latitude_reference)), lon_lat[:, 1] * LAT_KM]
    )


def consolidate_controls(stations: np.ndarray, fabdem: np.ndarray, copernicus: np.ndarray) -> list[dict]:
    groups: dict[tuple[float, float], list[int]] = defaultdict(list)
    for i, xy in enumerate(map(tuple, stations[:, :2])):
        groups[xy].append(i)

    controls: list[dict] = []
    for (lon, lat), indices in groups.items():
        heights = stations[indices, 2]
        station_h = float(np.median(heights))
        fab_h = float(np.median(fabdem[indices]))
        cop_h = float(np.median(copernicus[indices]))
        residual = station_h - fab_h
        reasons: list[str] = []
        if station_h == 0.0 and fab_h > 20.0:
            reasons.append("zero_altitude_incompatible_with_dem")
        if abs(residual) > OUTLIER_THRESHOLD_M:
            reasons.append("station_fabdem_residual_over_150m")
        if len(indices) > 1 and float(np.ptp(heights)) > 5.0:
            reasons.append("duplicate_coordinate_height_spread_over_5m")
        controls.append(
            {
                "longitude": lon,
                "latitude": lat,
                "station_altitude_m": station_h,
                "fabdem_altitude_m": fab_h,
                "copernicus_altitude_m": cop_h,
                "station_minus_fabdem_m": residual,
                "source_rows": [int(i + 2) for i in indices],
                "duplicate_height_spread_m": float(np.ptp(heights)),
                "accepted": not reasons,
                "exclusion_reasons": reasons,
            }
        )
    return controls


class RobustIDW:
    def __init__(self, xy: np.ndarray, residuals: np.ndarray):
        self.xy = xy
        self.targets = np.clip(residuals, -RESIDUAL_CLIP_M, RESIDUAL_CLIP_M)
        self.tree = cKDTree(xy)
        self.hull = Delaunay(xy)

    def predict(self, query_xy: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
        k = min(IDW_NEIGHBORS, len(self.xy))
        distances, indices = self.tree.query(query_xy, k=k)
        if k == 1:
            distances = distances[:, None]
            indices = indices[:, None]
        weights = 1.0 / (distances**IDW_POWER + IDW_SOFTENING_KM**IDW_POWER)
        correction = np.sum(weights * self.targets[indices], axis=1) / np.sum(weights, axis=1)
        inside = self.hull.find_simplex(query_xy) >= 0
        nearest_distance = distances[:, 0]
        support = np.ones(len(query_xy), dtype=float)
        support[~inside] = np.exp(-((nearest_distance[~inside] / OUTSIDE_TAPER_KM) ** 2))
        return correction, support


def metrics(errors: np.ndarray) -> dict[str, float]:
    return {
        "n": int(len(errors)),
        "bias_m": float(np.mean(errors)),
        "mae_m": float(np.mean(np.abs(errors))),
        "rmse_m": float(np.sqrt(np.mean(errors**2))),
        "median_error_m": float(np.median(errors)),
        "p95_absolute_error_m": float(np.quantile(np.abs(errors), 0.95)),
        "maximum_absolute_error_m": float(np.max(np.abs(errors))),
    }


def spatial_cross_validation(xy: np.ndarray, residuals: np.ndarray) -> dict:
    block_xy = np.floor((xy - xy.min(axis=0)) / BLOCK_SIZE_KM).astype(int)
    unique_blocks, inverse = np.unique(block_xy, axis=0, return_inverse=True)
    rng = np.random.default_rng(RANDOM_SEED)
    permutation = rng.permutation(len(unique_blocks))
    block_fold = np.empty(len(unique_blocks), dtype=int)
    block_fold[permutation] = np.arange(len(unique_blocks)) % CV_FOLDS
    folds = block_fold[inverse]
    corrected_errors = np.empty(len(residuals), dtype=float)
    for fold in range(CV_FOLDS):
        train = folds != fold
        test = ~train
        model = RobustIDW(xy[train], residuals[train])
        predicted, _ = model.predict(xy[test])
        corrected_errors[test] = predicted - residuals[test]
    return {
        "block_size_km": BLOCK_SIZE_KM,
        "folds": CV_FOLDS,
        "spatial_blocks": int(len(unique_blocks)),
        "fabdem_baseline": metrics(-residuals),
        "hybrid_cross_validated": metrics(corrected_errors),
    }


def copy_geotiff_tags(source: Image.Image) -> TiffImagePlugin.ImageFileDirectory_v2:
    tags = TiffImagePlugin.ImageFileDirectory_v2()
    for tag in (33550, 33922, 34735, 34736, 34737, 42112, 42113):
        if tag in source.tag_v2:
            tags[tag] = source.tag_v2[tag]
    return tags


def generate_tile(
    source_path: Path,
    destination_path: Path,
    lat0: int,
    lon0: int,
    model: RobustIDW,
    latitude_reference: float,
) -> dict:
    with Image.open(source_path) as source:
        base = np.asarray(source, dtype=np.float32)
        tags = copy_geotiff_tags(source)

    coarse_step = 1.0 / COARSE_SIZE
    padded_index = np.arange(-COARSE_PADDING, COARSE_SIZE + COARSE_PADDING)
    lons = lon0 + (padded_index + 0.5) * coarse_step
    lats = lat0 + 1.0 - (padded_index + 0.5) * coarse_step
    lon_grid, lat_grid = np.meshgrid(lons, lats)
    query_lon_lat = np.column_stack([lon_grid.ravel(), lat_grid.ravel()])
    query_xy = to_xy_km(query_lon_lat, latitude_reference)
    correction, support = model.predict(query_xy)
    padded_size = COARSE_SIZE + 2 * COARSE_PADDING
    coarse_correction = (correction * support).reshape(padded_size, padded_size).astype(np.float32)
    padded_correction = zoom(coarse_correction, COARSE_FACTOR, order=3, mode="nearest", grid_mode=True)
    crop = COARSE_PADDING * COARSE_FACTOR
    full_correction = padded_correction[crop:-crop, crop:-crop]
    full_correction = np.clip(full_correction, -RESIDUAL_CLIP_M, RESIDUAL_CLIP_M).astype(np.float32)
    hybrid = base + full_correction
    water = base <= 0.0
    hybrid[water] = base[water]
    hybrid[~water] = np.maximum(hybrid[~water], 0.0)

    destination_path.parent.mkdir(parents=True, exist_ok=True)
    output_image = Image.fromarray(hybrid, mode="F")
    output_image.save(destination_path, format="TIFF", compression="tiff_deflate", tiffinfo=tags)
    return {
        "tile": f"N{lat0:02d}E{lon0:03d}",
        "source": source_path.name,
        "output": destination_path.name,
        "minimum_m": float(np.nanmin(hybrid)),
        "maximum_m": float(np.nanmax(hybrid)),
        "correction_minimum_m": float(np.nanmin(full_correction)),
        "correction_maximum_m": float(np.nanmax(full_correction)),
        "changed_land_pixels": int(np.count_nonzero((~water) & (np.abs(full_correction) > 1e-5))),
    }


def sample_out1(stations: np.ndarray, path: Path) -> np.ndarray:
    grid = np.loadtxt(path)
    x = np.unique(grid[:, 0])
    y = np.unique(grid[:, 1])
    ordered = grid[np.lexsort((grid[:, 0], grid[:, 1]))]
    z = ordered[:, 2].reshape(len(y), len(x))
    interpolator = RegularGridInterpolator((y, x), z, method="linear", bounds_error=True)
    return interpolator(stations[:, [1, 0]])


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--stations", type=Path, required=True)
    parser.add_argument("--fabdem-dir", type=Path, required=True)
    parser.add_argument("--copernicus-dir", type=Path, required=True)
    parser.add_argument("--out1", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    stations = load_stations(args.stations)
    fabdem = sample_tiles(stations, args.fabdem_dir, "fabdem")
    copernicus = sample_tiles(stations, args.copernicus_dir, "copernicus")
    out1 = sample_out1(stations, args.out1)
    controls = consolidate_controls(stations, fabdem, copernicus)
    accepted = [item for item in controls if item["accepted"]]
    latitude_reference = float(np.mean([item["latitude"] for item in accepted]))
    accepted_lon_lat = np.asarray([[item["longitude"], item["latitude"]] for item in accepted])
    accepted_xy = to_xy_km(accepted_lon_lat, latitude_reference)
    residuals = np.asarray([item["station_minus_fabdem_m"] for item in accepted])
    model = RobustIDW(accepted_xy, residuals)
    cv = spatial_cross_validation(accepted_xy, residuals)

    tile_keys = sorted({(math.floor(lat), math.floor(lon)) for lon, lat in stations[:, :2]})
    tile_reports = []
    tile_dir = args.output_dir / "tiles"
    for lat0, lon0 in tile_keys:
        source = fabdem_path(args.fabdem_dir, lat0, lon0)
        destination = tile_dir / f"N{lat0:02d}E{lon0:03d}_HYBRID_FABDEM_STATIONS_30m.tif"
        tile_reports.append(generate_tile(source, destination, lat0, lon0, model, latitude_reference))

    # Model-space diagnostic at all original station coordinates.
    all_xy = to_xy_km(stations[:, :2], latitude_reference)
    predicted_correction, predicted_support = model.predict(all_xy)
    hybrid_station = fabdem + predicted_correction * predicted_support
    hybrid_station[fabdem <= 0.0] = fabdem[fabdem <= 0.0]
    positive = stations[:, 2] > 0.0
    accepted_source_rows = {row for item in accepted for row in item["source_rows"]}
    accepted_original = np.asarray([(i + 2) in accepted_source_rows for i in range(len(stations))])

    report = {
        "title": "MNT hybride FABDEM–stations pour les corrections gravimétriques au Cameroun",
        "station_count": int(len(stations)),
        "unique_coordinate_count": int(len(controls)),
        "accepted_control_count": int(len(accepted)),
        "excluded_control_count": int(len(controls) - len(accepted)),
        "input_extent_wgs84": {
            "west": float(stations[:, 0].min()),
            "east": float(stations[:, 0].max()),
            "south": float(stations[:, 1].min()),
            "north": float(stations[:, 1].max()),
        },
        "model": {
            "base_dem": "FABDEM v1.2",
            "output_resolution": "1 arc-second (~30 m)",
            "horizontal_crs": "WGS 84 / EPSG:4326",
            "correction_grid_resolution": "15 arc-seconds (~500 m), bicubic upsampling",
            "interpolator": "robust local inverse-distance weighting",
            "neighbors": IDW_NEIGHBORS,
            "power": IDW_POWER,
            "softening_km": IDW_SOFTENING_KM,
            "residual_clip_m": RESIDUAL_CLIP_M,
            "outside_convex_hull_taper_km": OUTSIDE_TAPER_KM,
            "outlier_threshold_m": OUTLIER_THRESHOLD_M,
        },
        "spatial_cross_validation": cv,
        "diagnostic_at_positive_altitude_stations": {
            "out1_surfer": metrics(out1[positive] - stations[positive, 2]),
            "fabdem": metrics(fabdem[positive] - stations[positive, 2]),
            "copernicus": metrics(copernicus[positive] - stations[positive, 2]),
            "hybrid_model_space_all": metrics(hybrid_station[positive] - stations[positive, 2]),
            "hybrid_model_space_accepted_controls": metrics(
                hybrid_station[positive & accepted_original] - stations[positive & accepted_original, 2]
            ),
        },
        "tiles": tile_reports,
        "excluded_controls": [item for item in controls if not item["accepted"]],
        "scientific_warning": (
            "Training-point diagnostics are not independent validation. The spatial block cross-validation "
            "is the principal model-selection result. Original station elevations remain the observation heights."
        ),
    }
    (args.output_dir / "hybrid_dem_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (args.output_dir / "excluded_controls.json").write_text(
        json.dumps(report["excluded_controls"], ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps({k: report[k] for k in ("accepted_control_count", "excluded_control_count", "spatial_cross_validation", "diagnostic_at_positive_altitude_stations")}, indent=2))


if __name__ == "__main__":
    main()
