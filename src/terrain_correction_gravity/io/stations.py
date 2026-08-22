"""Read and audit the original gravity-station workbook without modifying it."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


EXPECTED_HEADERS = ("Longitude", "Latitude", "Altitude", "A(air libre)")


@dataclass(frozen=True)
class Stations:
    longitude_deg: np.ndarray
    latitude_deg: np.ndarray
    altitude_m: np.ndarray
    free_air_anomaly_mgal: np.ndarray
    excel_row: np.ndarray

    def __len__(self) -> int:
        return int(self.longitude_deg.size)

    def as_matrix(self) -> np.ndarray:
        return np.column_stack(
            [
                self.longitude_deg,
                self.latitude_deg,
                self.altitude_m,
                self.free_air_anomaly_mgal,
            ]
        )


def load_stations(path: str | Path, sheet_name: str = "Feuil1") -> Stations:
    """Load the four original numeric columns and preserve Excel row numbers."""
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Missing worksheet {sheet_name!r}; found {workbook.sheetnames}")
    worksheet = workbook[sheet_name]
    headers = tuple(cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1)))
    if headers[:4] != EXPECTED_HEADERS:
        raise ValueError(f"Unexpected headers {headers[:4]!r}; expected {EXPECTED_HEADERS!r}")

    rows: list[list[float]] = []
    row_numbers: list[int] = []
    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        values = row[:4]
        if all(value is None for value in values):
            continue
        if any(value is None for value in values):
            raise ValueError(f"Incomplete measurement at Excel row {row_number}: {values!r}")
        try:
            rows.append([float(value) for value in values])
        except (TypeError, ValueError) as error:
            raise ValueError(f"Non-numeric measurement at Excel row {row_number}: {values!r}") from error
        row_numbers.append(row_number)

    data = np.asarray(rows, dtype=np.float64)
    if data.ndim != 2 or data.shape[1] != 4:
        raise ValueError("No valid station table was found")
    if not np.isfinite(data).all():
        raise ValueError("Station table contains NaN or infinite values")
    if np.any((data[:, 0] < -180) | (data[:, 0] > 180)):
        raise ValueError("Longitude outside [-180, 180] degrees")
    if np.any((data[:, 1] < -90) | (data[:, 1] > 90)):
        raise ValueError("Latitude outside [-90, 90] degrees")

    return Stations(
        longitude_deg=data[:, 0],
        latitude_deg=data[:, 1],
        altitude_m=data[:, 2],
        free_air_anomaly_mgal=data[:, 3],
        excel_row=np.asarray(row_numbers, dtype=np.int32),
    )


def audit_stations(stations: Stations) -> dict[str, object]:
    coordinates = np.column_stack([stations.longitude_deg, stations.latitude_deg])
    unique, counts = np.unique(coordinates, axis=0, return_counts=True)
    return {
        "count": len(stations),
        "unique_coordinate_count": int(unique.shape[0]),
        "duplicate_coordinate_count": int(np.sum(counts - 1)),
        "zero_altitude_count": int(np.count_nonzero(stations.altitude_m == 0.0)),
        "extent_wgs84": {
            "west": float(stations.longitude_deg.min()),
            "east": float(stations.longitude_deg.max()),
            "south": float(stations.latitude_deg.min()),
            "north": float(stations.latitude_deg.max()),
        },
        "altitude_m": {
            "minimum": float(stations.altitude_m.min()),
            "maximum": float(stations.altitude_m.max()),
        },
    }

